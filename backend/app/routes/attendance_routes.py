from flask import (
    Blueprint,
    request,
    jsonify,
    current_app,
    send_file,
)
from datetime import datetime, date, timedelta
from io import BytesIO

from flask_jwt_extended import (
    get_jwt,
    get_jwt_identity,
    verify_jwt_in_request,
)

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from app import db
from app.models.attendance import Attendance
from app.models.attendance_log import AttendanceLog
from app.models.employee import Employee
from app.models.shift import Shift
from app.models.branch import Branch
from app.models.department import Department
from app.models.designation import Designation
from app.services.gps_service import verify_employee_location
from app.utils.authorization import has_permission


attendance_bp = Blueprint(
    "attendance_bp",
    __name__
)


# ============================================================
# AUTHORIZATION HELPERS
# ============================================================

def get_current_user_context():
    """
    Returns the authenticated employee ID and role
    from the JWT.
    """

    verify_jwt_in_request()

    identity = get_jwt_identity()
    claims = get_jwt()

    role = str(
        claims.get("role", "")
    ).strip()

    return identity, role


def is_super_admin(role):
    return role.lower() in [
        "super admin",
        "super_admin",
        "admin"
    ]


def can_manage_attendance(role):
    """
    Super Admin always has Attendance management access.

    HR requires the database-controlled
    hr_attendance permission.
    """

    if is_super_admin(role):
        return True

    if role.lower() == "hr":
        return has_permission(
            role,
            "hr_attendance"
        )

    return False


def is_same_employee(
    current_employee_id,
    requested_employee_id
):
    return str(current_employee_id) == str(
        requested_employee_id
    )


# ============================================================
# ATTENDANCE MANAGEMENT FILTERS
# ============================================================

def apply_attendance_filters(
    query,
    employee_id=None,
    employee_code=None,
    branch_id=None,
    designation_id=None,
    status=None,
    date_from=None,
    date_to=None,
):
    """
    Apply management/reporting filters to an
    Attendance SQLAlchemy query.
    """

    if employee_id:
        query = query.filter(
            Attendance.employee_id == employee_id
        )

    if employee_code:
        employee = Employee.query.filter_by(
            employee_code=employee_code
        ).first()

        if employee:
            query = query.filter(
                Attendance.employee_id == employee.employee_id
            )
        else:
            query = query.filter(
                Attendance.employee_id == -1
            )

    if branch_id:
        employee_subquery = db.session.query(
            Employee.employee_id
        ).filter(
            Employee.branch_id == branch_id
        )

        query = query.filter(
            Attendance.employee_id.in_(employee_subquery)
        )

    if designation_id:
        employee_subquery = db.session.query(
            Employee.employee_id
        ).filter(
            Employee.designation_id == designation_id
        )

        query = query.filter(
            Attendance.employee_id.in_(employee_subquery)
        )

    if status:
        query = query.filter(
            Attendance.status == status
        )

    if date_from:
        query = query.filter(
            Attendance.attendance_date >= date_from
        )

    if date_to:
        query = query.filter(
            Attendance.attendance_date <= date_to
        )

    return query


def parse_report_date(value, field_name):
    """
    Parse YYYY-MM-DD date values used by
    attendance reporting filters.
    """

    if value is None or str(value).strip() == "":
        return None, None

    value = str(value).strip()

    try:
        return (
            datetime.strptime(
                value,
                "%Y-%m-%d"
            ).date(),
            None
        )

    except ValueError:

        return (
            None,
            f"{field_name} must be in YYYY-MM-DD format."
        )


def serialize_attendance(record):
    """
    Convert an Attendance record into the same
    JSON structure used by the frontend.
    """

    employee = Employee.query.get(
        record.employee_id
    )

    department = None
    designation = None
    branch = None
    shift = None

    if employee and employee.department_id:
        department = Department.query.get(
            employee.department_id
        )

    if employee and employee.designation_id:
        designation = Designation.query.get(
            employee.designation_id
        )

    if employee and employee.branch_id:
        branch = Branch.query.get(
            employee.branch_id
        )

    if employee and employee.shift_id:
        shift = Shift.query.get(
            employee.shift_id
        )

    return {

        "attendance_id":
            record.attendance_id,

        "employee_id":
            record.employee_id,

        "employee_name":
            employee.full_name
            if employee
            else None,

        "employee_code":
            employee.employee_code
            if employee
            else None,

        "department_name":
            department.department_name
            if department
            else None,

        "designation_name":
            designation.designation_name
            if designation
            else None,

        "branch_id":
            employee.branch_id
            if employee
            else None,

        "branch_name":
            branch.branch_name
            if branch
            else None,

        "role":
            employee.role
            if employee
            else None,

        "shift_name":
            shift.shift_name
            if shift
            else None,

        "attendance_date":
            str(record.attendance_date),

        "login_time":
            str(record.login_time)
            if record.login_time
            else None,

        "logout_time":
            str(record.logout_time)
            if record.logout_time
            else None,

        "working_seconds":
            record.working_seconds,

        "lunch_seconds":
            record.lunch_seconds,

        "lunch_start_time":
            str(record.lunch_start_time)
            if record.lunch_start_time
            else None,

        "lunch_end_time":
            str(record.lunch_end_time)
            if record.lunch_end_time
            else None,

        "status":
            record.status,
    }


# ============================================================
# EXCEL HELPERS
# ============================================================

EXPORT_COLUMNS = [
    "Attendance ID",
    "Employee ID",
    "Employee Code",
    "Employee Name",
    "Department",
    "Designation",
    "Branch",
    "Attendance Date",
    "Login Time",
    "Logout Time",
    "Lunch Start Time",
    "Lunch End Time",
    "Working Seconds",
    "Lunch Seconds",
    "Status",
]


IMPORT_COLUMNS = [
    "Employee Code",
    "Attendance Date",
    "Login Time",
    "Logout Time",
    "Lunch Start Time",
    "Lunch End Time",
    "Working Seconds",
    "Lunch Seconds",
    "Status",
]


def excel_datetime_value(value):
    """
    Convert Excel datetime/date/string values
    into Python values.
    """

    if value is None:
        return None

    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return value

    if isinstance(value, str):

        value = value.strip()

        if not value:
            return None

        formats = [
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%d-%m-%Y %H:%M:%S",
            "%d-%m-%Y %H:%M",
            "%Y-%m-%d",
            "%d-%m-%Y",
        ]

        for fmt in formats:

            try:
                return datetime.strptime(
                    value,
                    fmt
                )
            except ValueError:
                continue

    return None


def excel_time_value(value):
    """
    Convert Excel time/datetime/string into
    a Python datetime.time value.
    """

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.time()

    if hasattr(value, "hour") and hasattr(
        value,
        "minute"
    ):
        return value

    if isinstance(value, str):

        value = value.strip()

        if not value:
            return None

        formats = [
            "%H:%M:%S",
            "%H:%M",
            "%I:%M:%S %p",
            "%I:%M %p",
        ]

        for fmt in formats:

            try:
                return datetime.strptime(
                    value,
                    fmt
                ).time()
            except ValueError:
                continue

    return None


def parse_excel_attendance_date(value):
    """
    Convert Excel attendance date to date.
    """

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, str):

        value = value.strip()

        formats = [
            "%Y-%m-%d",
            "%d-%m-%Y",
            "%d/%m/%Y",
            "%Y/%m/%d",
        ]

        for fmt in formats:

            try:
                return datetime.strptime(
                    value,
                    fmt
                ).date()
            except ValueError:
                continue

    return None


def combine_date_time(
    attendance_date,
    value
):
    """
    Convert Excel date/time values into
    a Python datetime using attendance_date.
    """

    if value is None:
        return None

    # Excel datetime
    if isinstance(value, datetime):
        return value

    # Excel date
    if isinstance(value, date):
        return datetime.combine(
            attendance_date,
            value
        )

    # Excel time object
    if hasattr(value, "hour") and hasattr(
        value,
        "minute"
    ):
        return datetime.combine(
            attendance_date,
            value
        )

    # String values
    if isinstance(value, str):

        value = value.strip()

        if not value:
            return None

        # Full datetime strings
        datetime_formats = [
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%d-%m-%Y %H:%M:%S",
            "%d-%m-%Y %H:%M",
        ]

        for fmt in datetime_formats:
            try:
                return datetime.strptime(
                    value,
                    fmt
                )
            except ValueError:
                continue

        # Time-only strings
        time_formats = [
            "%H:%M:%S",
            "%H:%M",
            "%I:%M:%S %p",
            "%I:%M %p",
        ]

        for fmt in time_formats:
            try:
                parsed_time = datetime.strptime(
                    value,
                    fmt
                ).time()

                return datetime.combine(
                    attendance_date,
                    parsed_time
                )

            except ValueError:
                continue

    return None


    """
    Convert an Excel date/time value into
    a Python datetime using attendance_date.
    """

    if value is None:
        return None

    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return datetime.combine(
            attendance_date,
            value
        )

    parsed_time = excel_time_value(value)

    if parsed_time:
        return datetime.combine(
            attendance_date,
            parsed_time
        )

    return None


def format_excel_datetime(value):
    if not value:
        return None

    return value.strftime(
        "%Y-%m-%d %H:%M:%S"
    )


# ============================================================
# EMPLOYEE CHECK-IN
# ============================================================

@attendance_bp.route(
    "/api/attendance/login",
    methods=["POST"]
)
def employee_login():

    current_employee_id, role = (
        get_current_user_context()
    )

    print(
        "DEVELOPMENT_MODE =",
        current_app.config.get(
            "DEVELOPMENT_MODE"
        )
    )

    data = request.get_json() or {}

    employee_id = (
        data
        if isinstance(data, int)
        else data.get("employee_id")
    )

    latitude = (
        None
        if isinstance(data, int)
        else data.get("latitude")
    )

    longitude = (
        None
        if isinstance(data, int)
        else data.get("longitude")
    )

    if not employee_id:

        return jsonify({
            "success": False,
            "message":
                "employee_id is required."
        }), 400

    if not is_same_employee(
        current_employee_id,
        employee_id
    ):

        return jsonify({
            "success": False,
            "message":
                "Access denied. You can only record your own attendance."
        }), 403

    if latitude is None or longitude is None:

        return jsonify({
            "success": False,
            "message":
                "Current GPS location is required."
        }), 400

    try:

        latitude = float(latitude)
        longitude = float(longitude)

    except (TypeError, ValueError):

        return jsonify({
            "success": False,
            "message":
                "Invalid GPS coordinates."
        }), 400

    if not (-90 <= latitude <= 90):

        return jsonify({
            "success": False,
            "message":
                "Invalid latitude."
        }), 400

    if not (-180 <= longitude <= 180):

        return jsonify({
            "success": False,
            "message":
                "Invalid longitude."
        }), 400

    if current_app.config.get(
        "DEVELOPMENT_MODE",
        False
    ):

        gps_result = {
            "allowed": True,
            "branch_name": "Development Mode",
            "branch_id": None,
            "distance": 0,
            "allowed_radius": 0,
        }

    else:

        gps_result = verify_employee_location(
            employee_id=employee_id,
            latitude=latitude,
            longitude=longitude
        )

        if not gps_result["allowed"]:

            return jsonify(
                gps_result
            ), 403

    existing = Attendance.query.filter_by(
        employee_id=employee_id,
        attendance_date=date.today()
    ).first()

    if existing:

        return jsonify({
            "success": True,
            "message":
                "Attendance already recorded today",
            "already_logged_in": True
        }), 200

    employee = Employee.query.get(
        employee_id
    )

    if not employee:

        return jsonify({
            "success": False,
            "message":
                "Employee not found."
        }), 404

    status = "Working"

    if employee.shift_id:

        shift = Shift.query.get(
            employee.shift_id
        )

        if shift:

            now = datetime.now().time()

            grace_time = (
                datetime.combine(
                    date.today(),
                    shift.start_time
                )
                + timedelta(
                    minutes=shift.grace_minutes
                )
            ).time()

            if now > grace_time:
                status = "Late"
            else:
                status = "Present"

    attendance = Attendance(
        employee_id=employee_id,
        attendance_date=date.today(),
        login_time=datetime.now(),
        status=status,
        login_latitude=latitude,
        login_longitude=longitude,
        login_address=gps_result.get(
            "address"
        ),
        distance_from_branch=gps_result.get(
            "distance"
        ),
    )

    db.session.add(attendance)

    log = AttendanceLog(
        employee_id=employee_id,
        action="LOGIN",
        timestamp=datetime.now()
    )

    db.session.add(log)

    db.session.commit()

    return jsonify({
        "success": True,
        "message":
            "Login recorded successfully",
        "status": status,
        "gps": {
            "branch_name":
                gps_result["branch_name"],

            "branch_id":
                gps_result["branch_id"],

            "distance":
                gps_result["distance"],

            "allowed_radius":
                gps_result["allowed_radius"],
        }
    }), 200


# ============================================================
# LUNCH OUT
# ============================================================

@attendance_bp.route(
    "/api/attendance/lunch-out",
    methods=["POST"]
)
def lunch_out():

    current_employee_id, role = (
        get_current_user_context()
    )

    data = request.get_json() or {}

    employee_id = (
        data
        if isinstance(data, int)
        else data.get("employee_id")
    )

    if not employee_id:

        return jsonify({
            "success": False,
            "message":
                "employee_id is required."
        }), 400

    if not is_same_employee(
        current_employee_id,
        employee_id
    ):

        return jsonify({
            "success": False,
            "message":
                "Access denied. You can only manage your own attendance."
        }), 403

    attendance = Attendance.query.filter_by(
        employee_id=employee_id,
        attendance_date=date.today()
    ).first()

    if not attendance:

        return jsonify({
            "success": False,
            "message":
                "Attendance record not found"
        }), 404

    if attendance.logout_time:

        return jsonify({
            "success": False,
            "message":
                "Cannot start lunch after logout."
        }), 400

    if attendance.lunch_end_time:

        return jsonify({
            "success": False,
            "message":
                "Lunch break already used today"
        }), 400

    attendance.status = "Lunch Break"
    attendance.lunch_start_time = datetime.now()
    attendance.lunch_end_time = None

    log = AttendanceLog(
        employee_id=employee_id,
        action="LUNCH_OUT",
        timestamp=datetime.now()
    )

    db.session.add(log)
    db.session.commit()

    return jsonify({
        "success": True,
        "message":
            "Lunch break started",
        "attendance": serialize_attendance(
            attendance
        )
    }), 200


# ============================================================
# LUNCH IN
# ============================================================

@attendance_bp.route(
    "/api/attendance/lunch-in",
    methods=["POST"]
)
def lunch_in():

    current_employee_id, role = (
        get_current_user_context()
    )

    data = request.get_json() or {}

    employee_id = (
        data
        if isinstance(data, int)
        else data.get("employee_id")
    )

    if not employee_id:

        return jsonify({
            "success": False,
            "message":
                "employee_id is required."
        }), 400

    if not is_same_employee(
        current_employee_id,
        employee_id
    ):

        return jsonify({
            "success": False,
            "message":
                "Access denied. You can only manage your own attendance."
        }), 403

    attendance = Attendance.query.filter_by(
        employee_id=employee_id,
        attendance_date=date.today()
    ).first()

    if not attendance:

        return jsonify({
            "success": False,
            "message":
                "Attendance record not found"
        }), 404

    if attendance.logout_time:

        return jsonify({
            "success": False,
            "message":
                "Cannot return from lunch after logout."
        }), 400

    if not attendance.lunch_start_time:

        return jsonify({
            "success": False,
            "message":
                "Lunch break was not started"
        }), 400

    attendance.lunch_end_time = datetime.now()

    lunch_seconds = int(
        (
            attendance.lunch_end_time
            - attendance.lunch_start_time
        ).total_seconds()
    )

    attendance.lunch_seconds = (
        attendance.lunch_seconds or 0
    ) + lunch_seconds

    attendance.status = "Working"

    log = AttendanceLog(
        employee_id=employee_id,
        action="LUNCH_IN",
        timestamp=datetime.now()
    )

    db.session.add(log)
    db.session.commit()

    return jsonify({
        "success": True,
        "message":
            "Returned from lunch",
        "attendance": serialize_attendance(
            attendance
        )
    }), 200


# ============================================================
# CHECK OUT
# ============================================================

@attendance_bp.route(
    "/api/attendance/logout/<int:employee_id>",
    methods=["PUT"]
)
def employee_logout(employee_id):

    current_employee_id, role = (
        get_current_user_context()
    )

    if not is_same_employee(
        current_employee_id,
        employee_id
    ):

        return jsonify({
            "success": False,
            "message":
                "Access denied. You can only manage your own attendance."
        }), 403

    attendance = Attendance.query.filter_by(
        employee_id=employee_id,
        attendance_date=date.today()
    ).first()

    if not attendance:

        return jsonify({
            "success": False,
            "message":
                "Attendance record not found"
        }), 404

    if attendance.logout_time:

        return jsonify({
            "success": False,
            "message":
                "Already logged out"
        }), 400

    attendance.logout_time = datetime.now()

    total_duration = (
        attendance.logout_time
        - attendance.login_time
    )

    if (
        attendance.lunch_start_time
        and not attendance.lunch_end_time
    ):

        attendance.lunch_end_time = datetime.now()

        attendance.lunch_seconds = (
            attendance.lunch_seconds or 0
        ) + int(
            (
                attendance.lunch_end_time
                - attendance.lunch_start_time
            ).total_seconds()
        )

    working_seconds = max(
        0,
        int(
            total_duration.total_seconds()
        )
        - (attendance.lunch_seconds or 0)
    )

    attendance.working_seconds = (
        working_seconds
    )

    employee = Employee.query.get(
        employee_id
    )

    if employee and employee.shift_id:

        shift = Shift.query.get(
            employee.shift_id
        )

        if shift:

            logout_time = (
                attendance.logout_time.time()
            )

            if logout_time < shift.end_time:

                attendance.status = (
                    "Early Logout"
                )

            else:

                attendance.status = (
                    "Completed"
                )

    else:

        attendance.status = "Logged Out"

    log = AttendanceLog(
        employee_id=employee_id,
        action="LOGOUT",
        timestamp=datetime.now()
    )

    db.session.add(log)
    db.session.commit()

    return jsonify({
        "success": True,
        "message":
            "Logout recorded successfully",

        "working_seconds":
            attendance.working_seconds,

        "lunch_seconds":
            attendance.lunch_seconds,

        "status":
            attendance.status
    }), 200


# ============================================================
# TODAY'S PERSONAL ATTENDANCE
# ============================================================

@attendance_bp.route(
    "/api/attendance/today/<int:employee_id>",
    methods=["GET"]
)
def get_today_attendance(employee_id):

    current_employee_id, role = (
        get_current_user_context()
    )

    if not is_same_employee(
        current_employee_id,
        employee_id
    ):

        return jsonify({
            "success": False,
            "message":
                "Access denied. You can only view your own attendance."
        }), 403

    attendance = Attendance.query.filter_by(
        employee_id=employee_id,
        attendance_date=date.today()
    ).first()

    if not attendance:

        return jsonify({
            "success": True,
            "logged_in": False,
            "attendance": None
        }), 200

    employee = Employee.query.get(
        employee_id
    )

    branch = None

    if employee and employee.branch_id:

        branch = Branch.query.get(
            employee.branch_id
        )

    shift = None

    if employee and employee.shift_id:

        shift = Shift.query.get(
            employee.shift_id
        )

    logged_in = (
        attendance.logout_time is None
    )

    return jsonify({

        "success": True,

        "logged_in":
            logged_in,

        "attendance": {

            "attendance_id":
                attendance.attendance_id,

            "employee_id":
                attendance.employee_id,

            "employee_name":
                employee.full_name
                if employee
                else None,

            "employee_code":
                employee.employee_code
                if employee
                else None,

            "branch_id":
                employee.branch_id
                if employee
                else None,

            "branch_name":
                branch.branch_name
                if branch
                else None,

            "role":
                employee.role
                if employee
                else None,

            "shift_name":
                shift.shift_name
                if shift
                else None,

            "attendance_date":
                str(
                    attendance.attendance_date
                ),

            "login_time":
                str(
                    attendance.login_time
                )
                if attendance.login_time
                else None,

            "logout_time":
                str(
                    attendance.logout_time
                )
                if attendance.logout_time
                else None,

            "working_seconds":
                attendance.working_seconds,

            "lunch_seconds":
                attendance.lunch_seconds,

            "lunch_start_time":
                str(
                    attendance.lunch_start_time
                )
                if attendance.lunch_start_time
                else None,

            "lunch_end_time":
                str(
                    attendance.lunch_end_time
                )
                if attendance.lunch_end_time
                else None,

            "status":
                attendance.status
        }

    }), 200


# ============================================================
# ATTENDANCE MANAGEMENT / LIST
#
# SUPER ADMIN:
#     Full access
#
# HR:
#     Requires hr_attendance permission
#
# EMPLOYEE:
#     Own records only
# ============================================================

@attendance_bp.route(
    "/api/attendance",
    methods=["GET"]
)
def get_attendance():

    current_employee_id, role = (
        get_current_user_context()
    )

    requested_employee_id = (
        request.args.get(
            "employee_id",
            type=int
        )
    )

    employee_code = (
        request.args.get(
            "employee_code",
            type=str
        )
    )

    branch_id = (
    request.args.get(
        "branch_id",
        type=int
    )
)

    designation_id = (
    request.args.get(
        "designation_id",
        type=int
    )
)

    status = (  
        request.args.get(
            "status",
            type=str
        )
    )

    date_from, date_error = parse_report_date(
        request.args.get("date_from"),
        "date_from"
    )

    if date_error:
        return jsonify({
            "success": False,
            "message": date_error
        }), 400

    date_to, date_error = parse_report_date(
        request.args.get("date_to"),
        "date_to"
    )

    if date_error:
        return jsonify({
            "success": False,
            "message": date_error
        }), 400

    if date_from and date_to and date_from > date_to:

        return jsonify({
            "success": False,
            "message":
                "date_from cannot be later than date_to."
        }), 400

    # --------------------------------------------------------
    # Employee
    # --------------------------------------------------------

    if role.lower() == "employee":

        if requested_employee_id is None:

            return jsonify({
                "success": False,
                "message":
                    "employee_id is required."
            }), 403

        if not is_same_employee(
            current_employee_id,
            requested_employee_id
        ):

            return jsonify({
                "success": False,
                "message":
                    "Access denied. You can only view your own attendance."
            }), 403

        query = Attendance.query.filter(
            Attendance.employee_id
            == current_employee_id
        )

        # Employees cannot use management filters
        # to access another employee's records.
        query = apply_attendance_filters(
            query,
            employee_id=current_employee_id,
            date_from=date_from,
            date_to=date_to,
            status=status
        )

    # --------------------------------------------------------
    # Super Admin / HR
    # --------------------------------------------------------

    else:

        if not can_manage_attendance(role):

            return jsonify({
                "success": False,
                "message":
                    "Access denied. Attendance permission is required."
            }), 403

        query = Attendance.query

        query = apply_attendance_filters(
            query,
            employee_id=requested_employee_id,
            employee_code=employee_code,
            branch_id=branch_id,
            designation_id=designation_id,
            status=status,
            date_from=date_from,
            date_to=date_to
    )

    records = query.order_by(
        Attendance.attendance_date.desc(),
        Attendance.login_time.desc()
    ).all()

    attendance = [
        serialize_attendance(record)
        for record in records
    ]

    return jsonify({
        "success": True,
        "attendance": attendance,
        "count": len(attendance)
    }), 200


# ============================================================
# EXPORT ATTENDANCE TO EXCEL
#
# SUPER ADMIN:
#     Full access
#
# HR:
#     Requires hr_attendance permission
#
# EMPLOYEE:
#     Blocked
# ============================================================

@attendance_bp.route(
    "/api/attendance/export",
    methods=["GET"]
)
def export_attendance():

    current_employee_id, role = (
        get_current_user_context()
    )

    if not can_manage_attendance(role):

        return jsonify({
            "success": False,
            "message":
                "Access denied. Attendance permission is required."
        }), 403

    requested_employee_id = (
        request.args.get(
            "employee_id",
            type=int
        )
    )

    employee_code = (
        request.args.get(
            "employee_code",
            type=str
        )
    )

    branch_id = (
    request.args.get(
        "branch_id",
        type=int
    )
)

    designation_id = (
    request.args.get(
        "designation_id",
        type=int
    )
)

    status = (  
        request.args.get(
            "status",
            type=str
        )
    )

    date_from, date_error = parse_report_date(
        request.args.get("date_from"),
        "date_from"
    )

    if date_error:

        return jsonify({
            "success": False,
            "message": date_error
        }), 400

    date_to, date_error = parse_report_date(
        request.args.get("date_to"),
        "date_to"
    )

    if date_error:

        return jsonify({
            "success": False,
            "message": date_error
        }), 400

    if date_from and date_to and date_from > date_to:

        return jsonify({
            "success": False,
            "message":
                "date_from cannot be later than date_to."
        }), 400

    query = Attendance.query

    query = apply_attendance_filters(
        query,
        employee_id=requested_employee_id,
        employee_code=employee_code,
        branch_id=branch_id,
        designation_id=designation_id,
        status=status,
        date_from=date_from,
        date_to=date_to
    )

    records = query.order_by(
        Attendance.attendance_date.asc(),
        Attendance.employee_id.asc()
    ).all()

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Attendance Report"

    # --------------------------------------------------------
    # Header
    # --------------------------------------------------------

    for column_index, column_name in enumerate(
        EXPORT_COLUMNS,
        start=1
    ):

        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=column_name
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    row_number = 2

    for record in records:

        employee = Employee.query.get(
            record.employee_id
        )

        department = None
        designation = None
        branch = None

        if employee and employee.department_id:

            department = Department.query.get(
                employee.department_id
            )

        if employee and employee.designation_id:

            designation = Designation.query.get(
                employee.designation_id
            )

        if employee and employee.branch_id:

            branch = Branch.query.get(
                employee.branch_id
            )

        values = [
            record.attendance_id,

            record.employee_id,

            employee.employee_code
            if employee
            else None,

            employee.full_name
            if employee
            else None,

            department.department_name
            if department
            else None,

            designation.designation_name
            if designation
            else None,

            branch.branch_name
            if branch
            else None,

            record.attendance_date,

            record.login_time,

            record.logout_time,

            record.lunch_start_time,

            record.lunch_end_time,

            record.working_seconds,

            record.lunch_seconds,

            record.status,
        ]

        for column_index, value in enumerate(
            values,
            start=1
        ):

            worksheet.cell(
                row=row_number,
                column=column_index,
                value=value
            )

        row_number += 1

    # --------------------------------------------------------
    # Formatting
    # --------------------------------------------------------

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    for column_index in range(
        1,
        len(EXPORT_COLUMNS) + 1
    ):

        max_length = 0

        for cell in worksheet[
            get_column_letter(
                column_index
            )
        ]:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = min(
            max(max_length + 2, 12),
            35
        )

    # Date/time formatting

    for row in worksheet.iter_rows(
        min_row=2
    ):

        # Attendance Date
        row[7].number_format = (
            "yyyy-mm-dd"
        )

        # Login
        row[8].number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

        # Logout
        row[9].number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

        # Lunch Start
        row[10].number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

        # Lunch End
        row[11].number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

    # --------------------------------------------------------
    # Create file
    # --------------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    filename = (
        "attendance_report"
    )

    if date_from and date_to:

        filename += (
            f"_{date_from.isoformat()}"
            f"_to_{date_to.isoformat()}"
        )

    elif date_from:

        filename += (
            f"_from_{date_from.isoformat()}"
        )

    elif date_to:

        filename += (
            f"_until_{date_to.isoformat()}"
        )

    filename += ".xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# ============================================================
# IMPORT ATTENDANCE FROM EXCEL
#
# SUPER ADMIN:
#     Full access
#
# HR:
#     Requires hr_attendance permission
#
# EMPLOYEE:
#     Blocked
# ============================================================

@attendance_bp.route(
    "/api/attendance/import",
    methods=["POST"]
)
def import_attendance():

    current_employee_id, role = (
        get_current_user_context()
    )

    if not can_manage_attendance(role):

        return jsonify({
            "success": False,
            "message":
                "Access denied. Attendance permission is required."
        }), 403

    if "file" not in request.files:

        return jsonify({
            "success": False,
            "message":
                "Attendance Excel file is required."
        }), 400

    uploaded_file = request.files["file"]

    if not uploaded_file.filename:

        return jsonify({
            "success": False,
            "message":
                "No file selected."
        }), 400

    filename = uploaded_file.filename.lower()

    if not (
        filename.endswith(".xlsx")
        or filename.endswith(".xlsm")
    ):

        return jsonify({
            "success": False,
            "message":
                "Only .xlsx or .xlsm files are supported."
        }), 400

    try:

        workbook = load_workbook(
            uploaded_file,
            data_only=True
        )

    except Exception:

        return jsonify({
            "success": False,
            "message":
                "Unable to read the Excel file."
        }), 400

    worksheet = workbook.active

    # --------------------------------------------------------
    # Read header
    # --------------------------------------------------------

    header_values = []

    for cell in worksheet[1]:

        value = cell.value

        if value is None:
            header_values.append("")
        else:
            header_values.append(
                str(value).strip()
            )

    header_map = {
        name.lower(): index
        for index, name in enumerate(
            header_values
        )
        if name
    }

    missing_columns = [
        column
        for column in IMPORT_COLUMNS
        if column.lower() not in header_map
    ]

    if missing_columns:

        return jsonify({
            "success": False,
            "message":
                "Invalid attendance Excel format.",
            "missing_columns":
                missing_columns
        }), 400

    # --------------------------------------------------------
    # Process rows
    # --------------------------------------------------------

    errors = []
    processed = 0
    created = 0
    updated = 0

    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        row_is_empty = all(
            worksheet.cell(
                row=row_number,
                column=column_number
            ).value is None
            for column_number in range(
                1,
                worksheet.max_column + 1
            )
        )

        if row_is_empty:
            continue

        def get_value(column_name):

            index = header_map[
                column_name.lower()
            ]

            return worksheet.cell(
                row=row_number,
                column=index + 1
            ).value

        employee_code = get_value(
            "Employee Code"
        )

        attendance_date_value = get_value(
            "Attendance Date"
        )

        login_time_value = get_value(
            "Login Time"
        )

        logout_time_value = get_value(
            "Logout Time"
        )

        lunch_start_value = get_value(
            "Lunch Start Time"
        )

        lunch_end_value = get_value(
            "Lunch End Time"
        )

        working_seconds_value = get_value(
            "Working Seconds"
        )

        lunch_seconds_value = get_value(
            "Lunch Seconds"
        )

        status_value = get_value(
            "Status"
        )

        # ----------------------------------------------------
        # Required employee
        # ----------------------------------------------------

        if employee_code is None or str(
            employee_code
        ).strip() == "":

            errors.append({
                "row": row_number,
                "message":
                    "Employee Code is required."
            })

            continue

        employee_code = str(
            employee_code
        ).strip()

        employee = Employee.query.filter_by(
            employee_code=employee_code
        ).first()

        if not employee:

            errors.append({
                "row": row_number,
                "employee_code":
                    employee_code,
                "message":
                    "Employee not found."
            })

            continue

        # ----------------------------------------------------
        # Attendance date
        # ----------------------------------------------------

        attendance_date = (
            parse_excel_attendance_date(
                attendance_date_value
            )
        )

        if not attendance_date:

            errors.append({
                "row": row_number,
                "employee_code":
                    employee_code,
                "message":
                    "Invalid Attendance Date."
            })

            continue

        # ----------------------------------------------------
        # Login time
        # ----------------------------------------------------

        login_time = combine_date_time(
            attendance_date,
            login_time_value
        )

        logout_time = combine_date_time(
            attendance_date,
            logout_time_value
        )

        lunch_start_time = combine_date_time(
            attendance_date,
            lunch_start_value
        )

        lunch_end_time = combine_date_time(
            attendance_date,
            lunch_end_value
        )

        if login_time_value is not None and not login_time:

            errors.append({
                "row": row_number,
                "employee_code":
                    employee_code,
                "message":
                    "Invalid Login Time."
            })

            continue

        if logout_time_value is not None and not logout_time:

            errors.append({
                "row": row_number,
                "employee_code":
                    employee_code,
                "message":
                    "Invalid Logout Time."
            })

            continue

        if (
            lunch_start_value is not None
            and not lunch_start_time
        ):

            errors.append({
                "row": row_number,
                "employee_code":
                    employee_code,
                "message":
                    "Invalid Lunch Start Time."
            })

            continue

        if (
            lunch_end_value is not None
            and not lunch_end_time
        ):

            errors.append({
                "row": row_number,
                "employee_code":
                    employee_code,
                "message":
                    "Invalid Lunch End Time."
            })

            continue

        if (
            login_time
            and logout_time
            and logout_time < login_time
        ):

            errors.append({
                "row": row_number,
                "employee_code":
                    employee_code,
                "message":
                    "Logout Time cannot be before Login Time."
            })

            continue

        if (
            lunch_start_time
            and lunch_end_time
            and lunch_end_time < lunch_start_time
        ):

            errors.append({
                "row": row_number,
                "employee_code":
                    employee_code,
                "message":
                    "Lunch End Time cannot be before Lunch Start Time."
            })

            continue

        # ----------------------------------------------------
        # Numeric fields
        # ----------------------------------------------------

        try:

            working_seconds = (
                0
                if working_seconds_value is None
                or str(
                    working_seconds_value
                ).strip() == ""
                else int(
                    float(
                        working_seconds_value
                    )
                )
            )

            lunch_seconds = (
                0
                if lunch_seconds_value is None
                or str(
                    lunch_seconds_value
                ).strip() == ""
                else int(
                    float(
                        lunch_seconds_value
                    )
                )
            )

        except (
            TypeError,
            ValueError
        ):

            errors.append({
                "row": row_number,
                "employee_code":
                    employee_code,
                "message":
                    "Working Seconds and Lunch Seconds must be numeric."
            })

            continue

        if working_seconds < 0:

            errors.append({
                "row": row_number,
                "employee_code":
                    employee_code,
                "message":
                    "Working Seconds cannot be negative."
            })

            continue

        if lunch_seconds < 0:

            errors.append({
                "row": row_number,
                "employee_code":
                    employee_code,
                "message":
                    "Lunch Seconds cannot be negative."
            })

            continue

        status = (
            str(status_value).strip()
            if status_value is not None
            and str(status_value).strip()
            else "Present"
        )

        # ----------------------------------------------------
        # Find existing record
        # ----------------------------------------------------

        attendance = Attendance.query.filter_by(
            employee_id=employee.employee_id,
            attendance_date=attendance_date
        ).first()

        if attendance:

            attendance.login_time = (
                login_time
            )

            attendance.logout_time = (
                logout_time
            )

            attendance.lunch_start_time = (
                lunch_start_time
            )

            attendance.lunch_end_time = (
                lunch_end_time
            )

            attendance.working_seconds = (
                working_seconds
            )

            attendance.lunch_seconds = (
                lunch_seconds
            )

            attendance.status = status

            updated += 1

        else:

            attendance = Attendance(
                employee_id=employee.employee_id,
                attendance_date=attendance_date,
                login_time=login_time,
                logout_time=logout_time,
                lunch_start_time=lunch_start_time,
                lunch_end_time=lunch_end_time,
                working_seconds=working_seconds,
                lunch_seconds=lunch_seconds,
                status=status
            )

            db.session.add(attendance)

            created += 1

        processed += 1

    # --------------------------------------------------------
    # Safety:
    # If validation errors exist, do not commit partial data.
    # --------------------------------------------------------

    if errors:

        db.session.rollback()

        return jsonify({
            "success": False,
            "message":
                "Attendance import failed validation. No records were changed.",
            "processed": 0,
            "created": 0,
            "updated": 0,
            "errors": errors
        }), 400

    try:

        db.session.commit()

    except Exception as exc:

        db.session.rollback()

        current_app.logger.exception(
            "Attendance Excel import failed"
        )

        return jsonify({
            "success": False,
            "message":
                "Attendance import failed while saving records."
        }), 500

    return jsonify({
        "success": True,
        "message":
            "Attendance Excel imported successfully.",
        "processed":
            processed,
        "created":
            created,
        "updated":
            updated,
        "errors":
            []
    }), 200